"""
Report generation: Excel and PDF export for a forecast (+ optional
recommendation), and a multi-address comparison Excel workbook.
Built on openpyxl and reportlab -- both pure-Python, no system deps.
"""
import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def build_excel_report(feeder: str, address: str, forecast_df: pd.DataFrame,
                        recommendation=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"

    bold = Font(bold=True)
    header_fill = PatternFill(start_color="FFDCE6F1", end_color="FFDCE6F1", fill_type="solid")

    ws["A1"] = "Electricity Availability Forecast Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Feeder: {feeder}"
    ws["A3"] = f"Address: {address}"
    ws["A4"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    for r in (2, 3, 4):
        ws[f"A{r}"].font = bold

    start_row = 6
    ws.cell(row=start_row, column=1, value="Forecast").font = Font(bold=True, size=12)
    header_row = start_row + 1
    headers = list(forecast_df.columns)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = bold
        cell.fill = header_fill
    for r, row in enumerate(forecast_df.itertuples(index=False), start=header_row + 1):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 20

    if recommendation is not None:
        rec_start = header_row + len(forecast_df) + 3
        ws.cell(row=rec_start, column=1, value="Recommendation").font = Font(bold=True, size=12)
        fields = [
            ("Recommended solution", recommendation.primary.replace("_", " + ").title()),
            ("Daily need (kWh)", recommendation.daily_need_kwh),
            ("Avg. shortfall (h/day)", recommendation.avg_shortfall_hours),
            ("Solar size (kW)", recommendation.solar_kw),
            ("Solar cost (NGN)", recommendation.solar_cost_naira),
            ("Battery size (kWh)", recommendation.battery_kwh),
            ("Battery cost (NGN)", recommendation.battery_cost_naira),
            ("Generator size (kVA)", recommendation.generator_kva),
            ("Generator cost (NGN)", recommendation.generator_cost_naira),
            ("Generator fuel (NGN/day)", recommendation.generator_fuel_naira_per_day),
            ("Estimated payback (years)", recommendation.payback_years),
        ]
        for i, (label, val) in enumerate(fields):
            if val is None:
                continue
            r = rec_start + 1 + i
            ws.cell(row=r, column=1, value=label).font = bold
            ws.cell(row=r, column=2, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_comparison_excel(comparison_df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"

    bold = Font(bold=True)
    header_fill = PatternFill(start_color="FFDCE6F1", end_color="FFDCE6F1", fill_type="solid")

    ws["A1"] = "Address Comparison Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"

    header_row = 4
    for c, h in enumerate(comparison_df.columns, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = bold
        cell.fill = header_fill
    for r, row in enumerate(comparison_df.itertuples(index=False), start=header_row + 1):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    for c in range(1, len(comparison_df.columns) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf_report(feeder: str, address: str, forecast_df: pd.DataFrame,
                      recommendation=None) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleX", parent=styles["Title"], fontSize=16)
    story = []

    story.append(Paragraph("Electricity Availability Forecast Report", title_style))
    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph(f"<b>Feeder:</b> {feeder}", styles["Normal"]))
    story.append(Paragraph(f"<b>Address:</b> {address}", styles["Normal"]))
    story.append(Paragraph(f"<b>Generated:</b> {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]))
    story.append(Spacer(1, 0.6 * cm))

    story.append(Paragraph("Forecast", styles["Heading2"]))
    table_data = [list(forecast_df.columns)] + forecast_df.astype(str).values.tolist()
    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.6 * cm))

    if recommendation is not None:
        story.append(Paragraph("Recommendation", styles["Heading2"]))
        rec_rows = [
            ["Recommended solution", recommendation.primary.replace("_", " + ").title()],
            ["Daily need (kWh)", str(recommendation.daily_need_kwh)],
            ["Avg. shortfall (h/day)", str(recommendation.avg_shortfall_hours)],
        ]
        if recommendation.solar_kw:
            rec_rows += [
                ["Solar size (kW)", str(recommendation.solar_kw)],
                ["Solar cost (NGN)", f"{recommendation.solar_cost_naira:,.0f}"],
            ]
        if recommendation.battery_kwh:
            rec_rows += [
                ["Battery size (kWh)", str(recommendation.battery_kwh)],
                ["Battery cost (NGN)", f"{recommendation.battery_cost_naira:,.0f}"],
            ]
        if recommendation.generator_kva:
            rec_rows += [
                ["Generator size (kVA)", str(recommendation.generator_kva)],
                ["Generator cost (NGN)", f"{recommendation.generator_cost_naira:,.0f}"],
                ["Generator fuel (NGN/day)", f"{recommendation.generator_fuel_naira_per_day:,.0f}"],
            ]
        if recommendation.payback_years:
            rec_rows.append(["Estimated payback (years)", str(recommendation.payback_years)])

        t2 = Table(rec_rows)
        t2.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(t2)
        story.append(Spacer(1, 0.4 * cm))
        story.append(Paragraph(
            "Note: cost figures are indicative planning estimates using placeholder "
            "market rates, not live installer quotes.", styles["Italic"],
        ))

    doc.build(story)
    return buf.getvalue()
